import json
import re
from io import BytesIO

import frappe
import frappe.desk.query_report
import openpyxl
from frappe import _
from frappe.desk.query_report import (
    _export_query as original__export_query,
    build_xlsx_data as original_build_xlsx_data,
    clean_params,
    parse_json,
    run as original_run,
)
from frappe.utils import cint, cstr, formatdate
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# BSP report colours -- the same as on screen (public/css/bsp_report.bundle.css)
# and in report PDFs/prints (public/js/report_styling.js).
NAVY = "1F3A5F"
NAVY_BORDER = "34537A"
GRID_BORDER = "C8D2DE"
BAND = "F4F7FB"
TOTAL = "DCE6F2"
FILTER_LABEL = "E6ECF3"
FILTERS_PER_ROW = 4
DATE_RE = re.compile(r"^\d{4}-\d{2}-\d{2}$")

# Banding and borders touch every cell; past this many cells (e.g. a year of
# Stock Ledger) only the title, header and total row are styled so the export
# stays quick.
MAX_STYLED_CELLS = 300_000


@frappe.whitelist()
def export_query():
    """
    Override of frappe.desk.query_report.export_query. Same flow as Frappe's,
    but the export applies the user's column visibility preferences and Excel
    files get the BSP report styling -- for background (emailed) exports too.
    """
    from frappe.desk.utils import pop_csv_params

    form_params = frappe._dict(frappe.local.form_dict)
    csv_params = pop_csv_params(form_params)
    clean_params(form_params)
    parse_json(form_params)
    report_name = form_params.report_name
    frappe.permissions.can_export(
        frappe.get_cached_value("Report", report_name, "ref_doctype"),
        raise_exception=True,
    )

    if int(form_params.export_in_background or 0):
        user_email = frappe.get_cached_value("User", frappe.session.user, "email")
        frappe.enqueue(
            "bsp_engineering.overrides.query_report.run_export_query_job",
            user_email=user_email,
            form_params=form_params,
            csv_params=csv_params,
            queue="long",
            now=frappe.flags.in_test,
        )
        frappe.msgprint(
            _(
                "Your report is being generated in the background. You will receive an email on {0} with a download link once it is ready."
            ).format(user_email)
        )
        return

    return _export_query(form_params, csv_params)


def run_export_query_job(user_email, form_params, csv_params):
    from frappe.desk.utils import send_report_email

    report_name, file_extension, content = _export_query(form_params, csv_params, populate_response=False)
    send_report_email(
        user_email, report_name, file_extension, content, attached_to_name=form_params.report_name
    )


def _export_query(form_params, csv_params, populate_response=True):
    from frappe.desk.utils import provide_binary_file

    layout = {}

    def build_xlsx_data(data, visible_idx, include_indentation, include_filters=False, **kwargs):
        result, column_widths = original_build_xlsx_data(
            data, visible_idx, include_indentation, include_filters=include_filters, **kwargs
        )
        # Frappe puts the applied filters first, ended by an empty row, then
        # the column labels.
        has_filters = cint(include_filters) and data.filters
        layout["header_idx"] = result.index([]) + 1 if has_filters else 0
        layout["column_count"] = len(column_widths)
        # Frappe appends the total row as a plain list, which build_xlsx_data
        # passes through as-is -- so it is the last exported row only if the
        # very same object ends the export.
        layout["has_total_row"] = bool(
            cint(data.get("add_total_row")) and data.result and result and result[-1] is data.result[-1]
        )
        return result, column_widths

    try:
        # Temporarily monkey-patch so Frappe's _export_query uses our wrappers
        frappe.desk.query_report.run = _run_with_column_prefs
        frappe.desk.query_report.build_xlsx_data = build_xlsx_data
        exported = original__export_query(form_params, csv_params, populate_response=False)
    finally:
        # Ensure we restore the original methods
        frappe.desk.query_report.run = original_run
        frappe.desk.query_report.build_xlsx_data = original_build_xlsx_data

    if not exported:
        # No data: Frappe has already set up the "No data to export" page.
        return

    report_name, file_extension, content = exported
    if file_extension == "xlsx" and layout:
        try:
            filters = frappe.parse_json(form_params.applied_filters or "{}") or {}
            content = style_report_xlsx(
                content,
                _get_company(filters),
                _(form_params.report_name),
                _filter_items(filters),
                layout["header_idx"],
                layout["column_count"],
                layout["has_total_row"],
            )
        except Exception:
            # Unstyled is better than no export.
            frappe.log_error("BSP: could not style report Excel export")

    if not populate_response:
        return report_name, file_extension, content

    provide_binary_file(_(report_name), file_extension, content)


def _run_with_column_prefs(*args, **kwargs):
    """Frappe's run, with columns the user hid via Pick Columns marked hidden."""
    res = original_run(*args, **kwargs)

    report_name = kwargs.get("report_name") or (args[0] if args else None)

    if report_name and res and res.get("columns"):
        prefs = frappe.db.get_value(
            "Report Column Visibility",
            {"user": frappe.session.user, "report_name": report_name},
            "visible_columns",
        )

        if prefs:
            try:
                visible_columns = json.loads(prefs)
                for col in res.get("columns"):
                    col_id = col.get("id") or col.get("fieldname")
                    if col_id not in visible_columns:
                        col["hidden"] = 1
            except Exception as e:
                frappe.log_error(f"Error applying column visibility for export: {e}")

    return res


def _get_company(filters):
    return (
        filters.get("company")
        or frappe.defaults.get_user_default("Company")
        or frappe.db.get_single_value("Global Defaults", "default_company")
        or ""
    )


def _filter_items(filters):
    """[(label, value)] for the applied filters that have a value, Company
    left out (it heads the banner). The export only gets fieldnames, so the
    label is the fieldname spelled out ("from_date" -> "From Date")."""
    items = []
    for key, value in filters.items():
        if key == "company" or value in (None, "", [], 0):
            continue
        if isinstance(value, list | tuple):
            value = ", ".join(_filter_value(v) for v in value if v not in (None, ""))
        else:
            value = _filter_value(value)
        if value:
            items.append((_(frappe.unscrub(key)), value))
    return items


def _filter_value(value):
    if isinstance(value, str) and DATE_RE.match(value):
        return formatdate(value)
    return cstr(value)


def style_report_xlsx(
    content, company, title, filter_items, header_idx, column_count, has_total_row=False
):
    """Give a Frappe report export the BSP look, like the PDF: a navy banner
    with the company and report name, the applied filters in a strip below
    it, a centred white-on-navy header row frozen in place, banded rows and a
    tinted total row. `header_idx` is the 0-based row of the column labels in
    Frappe's export (after its own filter rows, which are replaced)."""
    wb = openpyxl.load_workbook(BytesIO(content))
    ws = wb.worksheets[0]
    width = max(column_count, 2)

    if header_idx:
        ws.delete_rows(1, header_idx)

    # Filters as label/value pairs, FILTERS_PER_ROW to a row; each pair gets
    # an equal block of columns, the value merged across the block's rest.
    per_row = min(FILTERS_PER_ROW, max(1, width // 2))
    block = width // per_row
    filter_rows = [filter_items[i : i + per_row] for i in range(0, len(filter_items), per_row)]

    banner_rows = 2 if company else 1
    top_rows = banner_rows + len(filter_rows) + 1  # + a blank row before the table
    ws.insert_rows(1, top_rows)
    header_row = top_rows + 1

    navy = PatternFill("solid", fgColor=NAVY)
    banner = [(company.upper(), 16, 30)] if company else []
    banner.append((title.upper(), 12, 22))
    for row, (text, size, height) in enumerate(banner, start=1):
        ws.cell(row=row, column=1, value=text)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
        cell = ws.cell(row=row, column=1)
        cell.font = Font(name="Calibri", size=size, bold=True, color="FFFFFF")
        cell.fill = navy
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = height

    grid_side = Side(style="thin", color=GRID_BORDER)
    grid = Border(left=grid_side, right=grid_side, top=grid_side, bottom=grid_side)
    label_fill = PatternFill("solid", fgColor=FILTER_LABEL)
    for offset, pairs in enumerate(filter_rows):
        row = banner_rows + 1 + offset
        for i, (label, value) in enumerate(pairs):
            first = i * block + 1
            # Title Case for filter labels (e.g. Total Purchase), not ALL CAPS
            label_cell = ws.cell(row=row, column=first, value=label)
            label_cell.font = Font(name="Calibri", size=9, bold=True, color="4A5B70")
            label_cell.fill = label_fill
            label_cell.border = grid
            label_cell.alignment = Alignment(vertical="center")
            ws.cell(row=row, column=first + 1, value=value)
            last = first + block - 1
            if last > first + 1:
                ws.merge_cells(start_row=row, start_column=first + 1, end_row=row, end_column=last)
            for col in range(first + 1, last + 1):
                ws.cell(row=row, column=col).border = grid
            value_cell = ws.cell(row=row, column=first + 1)
            value_cell.font = Font(name="Calibri", bold=True)
            value_cell.alignment = Alignment(vertical="center")

    navy_side = Side(style="thin", color=NAVY_BORDER)
    for col in range(1, column_count + 1):
        cell = ws.cell(row=header_row, column=col)
        cell.font = Font(name="Calibri", bold=True, color="FFFFFF")
        cell.fill = navy
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = Border(left=navy_side, right=navy_side, top=navy_side, bottom=navy_side)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)

    first_data_row = header_row + 1
    last_data_row = ws.max_row
    if column_count and (last_data_row - header_row) * column_count <= MAX_STYLED_CELLS:
        band = PatternFill("solid", fgColor=BAND)
        for row in ws.iter_rows(
            min_row=first_data_row, max_row=last_data_row, min_col=1, max_col=column_count
        ):
            banded = (row[0].row - first_data_row) % 2 == 1
            for cell in row:
                cell.border = grid
                if banded:
                    cell.fill = band

    if has_total_row and last_data_row >= first_data_row:
        for col in range(1, column_count + 1):
            cell = ws.cell(row=last_data_row, column=col)
            cell.font = Font(name="Calibri", bold=True)
            cell.fill = PatternFill("solid", fgColor=TOTAL)

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
